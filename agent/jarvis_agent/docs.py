"""Documents, on this PC only: read PDFs, Word, Excel, PowerPoint, text and images; keep files in JARVIS's own
file store so a document is read once and asked about many times; and create Word, Excel, PowerPoint and PDF
files. Nothing here depends on an AI service: text comes out with local libraries (and Windows' own OCR for
images), and the AI only ever sees the passages that answer the question, each tagged with where it came from
(page, slide, sheet), so answers can quote the exact sentence they rely on.

The store lives in <data>/files: a copy of each file, its extracted text by page, and an index."""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import shutil
import uuid
from pathlib import Path
from typing import Any

from .actions import ActionError, _guard, _resolve, _str, action

TEXT = {".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".ts", ".html", ".xml", ".yaml", ".yml", ".ini", ".cfg"}
IMAGES = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
READABLE = TEXT | IMAGES | {".pdf", ".docx", ".xlsx", ".pptx"}
MAX_STORE_MB = 200
PASSAGE_CHARS = 700


def _store() -> Path:
    d = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "JARVIS" / "files"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _index() -> dict[str, Any]:
    f = _store() / "index.json"
    try:
        return json.loads(f.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _save_index(idx: dict[str, Any]) -> None:
    tmp = _store() / "index.tmp"
    tmp.write_text(json.dumps(idx, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, _store() / "index.json")


# ---------------- extracting text ----------------

def _pages(p: Path) -> list[tuple[str, str]]:
    """The document's text as (where, text) parts: ("p. 3", ...), ("slide 2", ...), ("sheet Sales", ...)."""
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(p))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:  # noqa: BLE001
                raise ActionError(f"{p.name} is password-protected.") from e
        return [(f"p. {i}", (page.extract_text() or "").strip()) for i, page in enumerate(reader.pages, 1)]
    if ext == ".docx":
        import docx

        d = docx.Document(str(p))
        parts, chunk, n = [], [], 1
        for para in d.paragraphs:
            if para.text.strip():
                chunk.append(para.text.strip())
            if sum(len(c) for c in chunk) > 2500:  # Word has no fixed pages: number sections instead
                parts.append((f"section {n}", "\n".join(chunk)))
                chunk, n = [], n + 1
        for t, table in enumerate(d.tables, 1):
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
            chunk.append(f"[table {t}]\n" + "\n".join(rows))
        if chunk:
            parts.append((f"section {n}", "\n".join(chunk)))
        return parts
    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                if any(v is not None for v in r):
                    rows.append(" | ".join("" if v is None else str(v) for v in r))
                if len(rows) >= 2000:
                    rows.append("…(more rows)")
                    break
            out.append((f"sheet {ws.title}", "\n".join(rows)))
        wb.close()  # read-only workbooks keep the file open until closed
        return out
    if ext == ".pptx":
        from pptx import Presentation

        out = []
        for i, slide in enumerate(Presentation(str(p)).slides, 1):
            texts = [sh.text_frame.text.strip() for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                texts.append("[notes] " + slide.notes_slide.notes_text_frame.text.strip())
            out.append((f"slide {i}", "\n".join(texts)))
        return out
    if ext in IMAGES:
        from .uictl import _OCR_SCRIPT, _ps

        lines = _ps(_OCR_SCRIPT, {"path": str(p)}, timeout=60).get("lines") or []
        return [("image", "\n".join(l["text"] for l in sorted(lines, key=lambda l: (l["y"] // 8, l["x"]))))]
    if ext in TEXT:
        text = p.read_text(encoding="utf-8", errors="replace")
        return [(f"part {i // 6000 + 1}", text[i:i + 6000]) for i in range(0, max(1, len(text)), 6000)]
    raise ActionError(f"I can't read {ext or 'that kind of'} files yet. I read PDF, Word, Excel, PowerPoint, text and images.")


def _words(text: str) -> set[str]:
    return {w for w in re.findall(r"[a-z0-9]+", text.lower()) if len(w) > 2}


def _passages(parts: list[tuple[str, str]], query: str, limit: int = 6) -> list[dict[str, str]]:
    """The passages that best answer a question, each with where it came from."""
    want = _words(query)
    chunks = []
    for where, text in parts:
        paras = [x.strip() for x in re.split(r"\n\s*\n|(?<=[.!?])\s{2,}", text) if x.strip()] or [text]
        buf = ""
        for para in paras:
            if len(buf) + len(para) > PASSAGE_CHARS and buf:
                chunks.append((where, buf))
                buf = ""
            buf = (buf + "\n" + para).strip()
        if buf:
            chunks.append((where, buf))
    scored = sorted(((len(want & _words(t)), i, where, t) for i, (where, t) in enumerate(chunks)), key=lambda x: (-x[0], x[1]))
    best = [x for x in scored if x[0] > 0][:limit] or scored[:limit]
    return [{"where": where, "text": t} for _s, _i, where, t in sorted(best, key=lambda x: x[1])]


def _page_range(spec: str, total: int) -> set[int]:
    out: set[int] = set()
    for piece in re.split(r"[,\s]+", spec.strip()):
        if "-" in piece:
            a, _, b = piece.partition("-")
            if a.isdigit() and b.isdigit():
                out |= set(range(int(a), min(int(b), total) + 1))
        elif piece.isdigit():
            out.add(int(piece))
    return out


# ---------------- the file store ----------------

def _stored(ref: str) -> tuple[Path, list[tuple[str, str]], dict[str, Any] | None]:
    """A file by store id or path: its location and text parts (from the store's cache when it is there)."""
    idx = _index()
    entry = idx.get(ref) or next((e for e in idx.values() if e["name"].lower() == ref.lower() or e.get("source", "").lower() == ref.lower()), None)
    if entry:
        cache = _store() / f"{entry['id']}.json"
        try:
            return _store() / entry["file"], [tuple(x) for x in json.loads(cache.read_text(encoding="utf-8"))], entry
        except (OSError, ValueError):
            pass
    p = _resolve(ref)
    if not p.is_file():
        raise ActionError(f"No file or stored file named '{ref}'.")
    return p, _pages(p), entry


@action
def store_file(args: dict[str, Any]) -> dict[str, Any]:
    """Keep a file in JARVIS's file store: copied once, text extracted once, then asked about any number of times."""
    src = _resolve(_str(args, "path"))
    if not src.is_file():
        raise ActionError(f"{src} is not a file.")
    if src.suffix.lower() not in READABLE:
        raise ActionError(f"I can't read {src.suffix} files, so storing it wouldn't help.")
    used = sum(f.stat().st_size for f in _store().iterdir() if f.is_file()) / 2**20
    if used + src.stat().st_size / 2**20 > MAX_STORE_MB:
        raise ActionError(f"The file store is full ({used:.0f} MB of {MAX_STORE_MB}); remove something first.")
    idx = _index()
    same = next((e for e in idx.values() if e.get("source") == str(src) and e.get("size") == src.stat().st_size), None)
    if same:
        return {"id": same["id"], "name": same["name"], "already_stored": True, "parts": same["parts"]}
    fid = uuid.uuid4().hex[:8]
    parts = _pages(src)
    shutil.copy2(src, _store() / f"{fid}{src.suffix.lower()}")
    (_store() / f"{fid}.json").write_text(json.dumps(parts, ensure_ascii=False), encoding="utf-8")
    idx[fid] = {"id": fid, "name": src.name, "source": str(src), "file": f"{fid}{src.suffix.lower()}", "size": src.stat().st_size,
                "parts": len(parts), "chars": sum(len(t) for _w, t in parts), "added": dt.datetime.now().strftime("%Y-%m-%d %H:%M")}
    _save_index(idx)
    empty = not any(t.strip() for _w, t in parts)
    return {"id": fid, "name": src.name, "parts": len(parts), **({"note": "No text found: it may be a scanned PDF."} if empty else {})}


@action
def stored_files(_: dict[str, Any]) -> dict[str, Any]:
    idx = _index()
    return {"files": [{k: e[k] for k in ("id", "name", "parts", "added", "source")} for e in sorted(idx.values(), key=lambda e: e["added"], reverse=True)]}


@action
def forget_file(args: dict[str, Any]) -> dict[str, Any]:
    fid = _str(args, "id")
    idx = _index()
    e = idx.pop(fid, None) or idx.pop(next((k for k, v in idx.items() if v["name"].lower() == fid.lower()), ""), None)
    if not e:
        raise ActionError(f"No stored file '{fid}'.")
    for f in (_store() / e["file"], _store() / f"{e['id']}.json"):
        f.unlink(missing_ok=True)
    _save_index(idx)
    return {"removed": e["name"]}


@action
def read_document(args: dict[str, Any]) -> dict[str, Any]:
    """Read a document (path or store id): the passages answering `question`, or the given pages, with locations."""
    p, parts, entry = _stored(_str(args, "file"))
    name = entry["name"] if entry else p.name
    if not any(t.strip() for _w, t in parts):
        raise ActionError(f"{name} has no readable text (a scanned PDF needs OCR first: open it and use read_window).")
    pages = _str(args, "pages", False)
    question = _str(args, "question", False)
    if pages:
        wanted = _page_range(pages, len(parts))
        chosen = [{"where": w, "text": t[:6000]} for i, (w, t) in enumerate(parts, 1) if i in wanted]
    elif question:
        chosen = _passages(parts, question)
    else:
        chosen, size = [], 0
        for w, t in parts:
            if size > 12000:
                break
            chosen.append({"where": w, "text": t[:6000]})
            size += len(t)
    return {"document": name, "total_parts": len(parts), "passages": chosen,
            "cite_as": f"Quote the exact words you rely on and say where they are, e.g. ({name}, {chosen[0]['where'] if chosen else 'p. 1'})."}


# ---------------- creating documents ----------------

def _blocks(content: str) -> list[tuple[str, Any]]:
    """Markdown-style text as blocks: ("h1".."h3", text), ("bullet", text), ("table", rows), ("p", text)."""
    out: list[tuple[str, Any]] = []
    table: list[list[str]] = []
    for raw in content.splitlines() + [""]:
        line = raw.rstrip()
        if line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                table.append(cells)
            continue
        if table:
            out.append(("table", table))
            table = []
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            out.append((f"h{len(m.group(1))}", m.group(2).strip()))
        elif re.match(r"^\s*([-*•]|\d+[.)])\s+", line):
            out.append(("bullet", re.sub(r"^\s*([-*•]|\d+[.)])\s+", "", line)))
        elif line.strip():
            out.append(("p", line.strip()))
    return out


def _target(path: str, title: str, ext: str) -> Path:
    if path:
        p = _resolve(path)
        if p.suffix.lower() != ext:
            p = p.with_suffix(ext)
    else:
        safe = re.sub(r'[<>:"/\\|?*]+', " ", title).strip()[:80] or "Document"
        p = Path.home() / "Documents" / "JARVIS" / f"{safe}{ext}"
    _guard(p)
    p.parent.mkdir(parents=True, exist_ok=True)
    n = 2
    base = p
    while p.exists():  # never overwrite
        p = base.with_name(f"{base.stem} ({n}){base.suffix}")
        n += 1
    return p


def _docx(p: Path, title: str, blocks: list[tuple[str, Any]]) -> None:
    import docx

    d = docx.Document()
    if title:
        d.add_heading(title, 0)
    for kind, v in blocks:
        if kind.startswith("h"):
            d.add_heading(v, int(kind[1]))
        elif kind == "bullet":
            d.add_paragraph(v, style="List Bullet")
        elif kind == "table":
            t = d.add_table(rows=len(v), cols=max(len(r) for r in v))
            t.style = "Table Grid"
            for i, row in enumerate(v):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = cell
        else:
            d.add_paragraph(v)
    d.save(str(p))


def _pdf(p: Path, title: str, blocks: list[tuple[str, Any]]) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    font = "Helvetica"
    arial = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "arial.ttf"
    if arial.is_file():  # a Unicode font, so names and symbols beyond Latin-1 print
        pdf.add_font("Arial", "", str(arial))
        bold = arial.with_name("arialbd.ttf")
        pdf.add_font("Arial", "B", str(bold if bold.is_file() else arial))
        font = "Arial"

    def text(s: str, size: int = 11, style: str = "", gap: float = 2) -> None:
        pdf.set_font(font, style, size)
        pdf.multi_cell(0, size * 0.5, s, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    if title:
        text(title, 20, "B", 4)
    for kind, v in blocks:
        if kind.startswith("h"):
            text(v, {"h1": 16, "h2": 14, "h3": 12}[kind], "B")
        elif kind == "bullet":
            text("•  " + v, gap=1)
        elif kind == "table":
            pdf.set_font(font, "", 10)
            with pdf.table() as t:
                for row in v:
                    r = t.row()
                    for cell in row:
                        r.cell(cell)
            pdf.ln(3)
        else:
            text(v)
    pdf.output(str(p))


def _pptx(p: Path, title: str, blocks: list[tuple[str, Any]]) -> int:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    slides: list[tuple[str, list[str]]] = []
    for kind, v in blocks:
        if kind in ("h1", "h2"):
            slides.append((v, []))
        elif kind == "table":
            slides[-1:] = slides[-1:] or [("", [])]
            slides[-1][1].extend(" | ".join(r) for r in v)
        else:
            if not slides:
                slides.append((title or "", []))
            slides[-1][1].append(v if kind != "h3" else v.upper())
    if title:
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = title
    for head, lines in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = head
        body = s.placeholders[1].text_frame
        body.clear()
        for i, line in enumerate(lines[:12]):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = line
            para.font.size = Pt(20 if len(lines) <= 6 else 16)
    prs.save(str(p))
    return len(prs.slides)


def _xlsx(p: Path, title: str, content: str) -> int:
    """Each '# Sheet name' starts a sheet; rows are markdown table lines or comma-separated lines."""
    import csv
    import io

    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet = None
    first_row = True
    for raw in content.splitlines():
        line = raw.strip()
        m = re.match(r"^#{1,3}\s+(.*)", line)
        if m or sheet is None:
            sheet = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", " ", (m.group(1) if m else title or "Sheet1"))[:31] or "Sheet")
            first_row = True
            if m:
                continue
        if not line or re.fullmatch(r"\|?(\s*:?-{2,}:?\s*\|)+\s*:?-*:?\s*\|?", line):
            continue
        cells = [c.strip() for c in (line.strip("|").split("|") if line.startswith("|") else next(csv.reader(io.StringIO(line))))]
        values: list[Any] = []
        for c in cells:
            try:
                values.append(float(c.replace(",", "")) if re.fullmatch(r"-?[\d,]*\.?\d+", c) else c)
            except ValueError:
                values.append(c)
        sheet.append(values)
        if first_row:
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
            first_row = False
    for ws in wb.worksheets:
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) + 2 for c in col)))
    wb.save(str(p))
    return len(wb.worksheets)


@action
def create_document(args: dict[str, Any]) -> dict[str, Any]:
    """Create a Word, Excel, PowerPoint or PDF file from markdown-style text. Never overwrites."""
    kind = (_str(args, "kind") or "").lower().lstrip(".")
    kind = {"word": "docx", "doc": "docx", "excel": "xlsx", "xls": "xlsx", "spreadsheet": "xlsx", "powerpoint": "pptx",
            "ppt": "pptx", "slides": "pptx", "presentation": "pptx"}.get(kind, kind)
    if kind not in ("docx", "xlsx", "pptx", "pdf"):
        raise ActionError("kind must be docx (Word), xlsx (Excel), pptx (PowerPoint) or pdf.")
    title = _str(args, "title", False) or ""
    content = str(args.get("content") or "")
    if not content.strip():
        raise ActionError("content is empty.")
    p = _target(_str(args, "path", False) or "", title, f".{kind}")
    extra: dict[str, Any] = {}
    if kind == "docx":
        _docx(p, title, _blocks(content))
    elif kind == "pdf":
        _pdf(p, title, _blocks(content))
    elif kind == "pptx":
        extra["slides"] = _pptx(p, title, _blocks(content))
    else:
        extra["sheets"] = _xlsx(p, title, content)
    if args.get("open"):
        os.startfile(str(p))
    return {"created": str(p), "kind": kind, "size_kb": round(p.stat().st_size / 1024, 1), **extra, "opened": bool(args.get("open"))}
